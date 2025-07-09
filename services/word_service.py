from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.shared import qn
import datetime
import os
import re


def set_cell_properties(cell, text, bold=False, align='LEFT', font_name='仿宋_GB2312', font_size=11):
    """一个基础的单元格/段落设置函数，处理文本、格式和换行"""
    # 检查传入的是单元格还是段落
    if isinstance(cell, type(document.add_paragraph())):
        p = cell
    else:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.text = ''
        p = cell.paragraphs[0]

    p.text = ''  # 清空段落现有内容

    lines = str(text).split('\n')

    # 添加第一行
    run = p.add_run(lines[0])

    # 设置字体
    font = run.font
    font.name = font_name
    font.size = Pt(font_size)
    font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

    # 设置对齐
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == 'CENTER' else WD_ALIGN_PARAGRAPH.LEFT

    # 添加后续行
    if len(lines) > 1:
        # 如果是单元格，就在单元格里加新段落
        if not isinstance(cell, type(document.add_paragraph())):
            for line in lines[1:]:
                p_new = cell.add_paragraph()
                run_new = p_new.add_run(line)
                font_new = run_new.font
                font_new.name = font_name
                font_new.size = Pt(font_size)
                font_new.bold = bold
                run_new._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
                p_new.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == 'CENTER' else WD_ALIGN_PARAGRAPH.LEFT
        else:  # 如果是段落，则通过添加换行符实现
            for line in lines[1:]:
                p.add_run('\n' + line)


def create_word_report(data: dict, image_input) -> str:
    """
    以3列表格为基础，严格按照用户提供的文档格式生成报告。
    """
    image_list = []
    if isinstance(image_input, str) and image_input:
        image_list = [{'label': '问题照片', 'path': image_input}]
    elif isinstance(image_input, list):
        image_list = image_input

    global document
    document = Document()
    sections = document.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # [cite_start]标题 [cite: 1]
    title_para = document.add_paragraph()
    set_cell_properties(title_para, '技术监督典型案例收集模板', bold=True, align='CENTER', font_name='黑体',
                        font_size=14)

    # 创建3列表格
    table = document.add_table(rows=30, cols=3)
    table.style = 'Table Grid'

    table.columns[0].width = Inches(1.5)
    table.columns[1].width = Inches(1.8)
    table.columns[2].width = Inches(4.0)

    row_idx = 0

    # 内联函数，用于添加合并了后两列的行，减少代码重复
    def add_merged_row(label, value_key):
        nonlocal row_idx
        set_cell_properties(table.cell(row_idx, 0), label, bold=True)
        merged_cell = table.cell(row_idx, 1);
        merged_cell.merge(table.cell(row_idx, 2))
        set_cell_properties(merged_cell, data.get(value_key, '待补充'))
        row_idx += 1

    # [cite_start]单位名称 [cite: 2]
    add_merged_row('*单位名称', 'unit_name')
    # [cite_start]案例名称 [cite: 2]
    add_merged_row('*案例名称', 'case_name')

    # [cite_start]工程信息 [cite: 2]
    start_row_project = row_idx
    cat_cell = table.cell(start_row_project, 0)
    cat_cell.merge(table.cell(start_row_project + 3, 0))
    set_cell_properties(cat_cell, '工程信息', bold=True, align='CENTER')

    project_info = data.get('project_info', {})
    set_cell_properties(table.cell(start_row_project, 1), '*工程名称')
    set_cell_properties(table.cell(start_row_project, 2), project_info.get('project_name', '待补充'))

    set_cell_properties(table.cell(start_row_project + 1, 1), '*站线名称')
    set_cell_properties(table.cell(start_row_project + 1, 2), project_info.get('station_name', '待补充'))

    set_cell_properties(table.cell(start_row_project + 2, 1), '*工程类型')
    set_cell_properties(table.cell(start_row_project + 2, 2), project_info.get('project_type', '待补充'))

    set_cell_properties(table.cell(start_row_project + 3, 1), '*工程电压等级')
    set_cell_properties(table.cell(start_row_project + 3, 2), project_info.get('project_voltage', '待补充'))
    row_idx += 4

    # [cite_start]设备信息 [cite: 2]
    start_row_device = row_idx
    cat_cell = table.cell(start_row_device, 0);
    cat_cell.merge(table.cell(start_row_device + 2, 0))
    set_cell_properties(cat_cell, '设备信息', bold=True, align='CENTER')

    device_info = data.get('device_info', {})
    set_cell_properties(table.cell(start_row_device, 1), '*设备类型')
    set_cell_properties(table.cell(start_row_device, 2), device_info.get('device_type', '待补充'))

    set_cell_properties(table.cell(start_row_device + 1, 1), '*设备电压等级')
    set_cell_properties(table.cell(start_row_device + 1, 2), device_info.get('device_voltage', '待补充'))

    set_cell_properties(table.cell(start_row_device + 2, 1), '*生产厂家全称（或责任单位）')
    set_cell_properties(table.cell(start_row_device + 2, 2), device_info.get('manufacturer', '待补充'))
    row_idx += 3

    # [cite_start]监督专业 & 监督阶段 [cite: 2]
    add_merged_row('*监督专业', 'supervision_major')
    add_merged_row('*监督阶段', 'supervision_stage')

    # [cite_start]依据细则（标准） [cite: 2]
    start_row_reg = row_idx
    cat_cell = table.cell(start_row_reg, 0);
    cat_cell.merge(table.cell(start_row_reg + 2, 0))
    set_cell_properties(cat_cell, '依据细则（标准）', bold=True, align='CENTER')

    regulation = data.get('regulation', {})
    set_cell_properties(table.cell(start_row_reg, 1), '*细则（标准）名称')
    set_cell_properties(table.cell(start_row_reg, 2), regulation.get('title', '待补充'))

    set_cell_properties(table.cell(start_row_reg + 1, 1), '*条款序号')
    set_cell_properties(table.cell(start_row_reg + 1, 2), regulation.get('clause', '待补充'))

    set_cell_properties(table.cell(start_row_reg + 2, 1), '*监督标准')
    # 使用规范中的监督依据作为监督标准
    supervision_standard = regulation.get('points', '待补充') if regulation.get('points') else '待补充'
    set_cell_properties(table.cell(start_row_reg + 2, 2), supervision_standard)
    row_idx += 3

    # [cite_start]发现时间 [cite: 2]
    add_merged_row('*发现时间', 'discovery_date')

    # [cite_start]问题描述, 原因分析, 等... [cite: 2]
    add_merged_row('*原始描述', 'description')
    add_merged_row('*问题描述', 'expanded_description')
    add_merged_row('*问题原因分析', 'analysis')
    add_merged_row('*监督意见', 'suggestions')
    add_merged_row('*实际整改措施', 'rectification_measures')
    add_merged_row('其他要说明的问题', 'other_issues')

    # [cite_start]问题照片 [cite: 2]
    set_cell_properties(table.cell(row_idx, 0), '*问题照片', bold=True)
    photo_cell_merged = table.cell(row_idx, 1);
    photo_cell_merged.merge(table.cell(row_idx, 2))

    if image_list:
        photo_cell_merged.text = ''
        nested_table = photo_cell_merged.add_table(rows=len(image_list), cols=2)
        nested_table.style = 'Table Grid'
        nested_table.columns[0].width = Inches(2.0)
        nested_table.columns[1].width = Inches(3.8)

        for i, img_info in enumerate(image_list):
            set_cell_properties(nested_table.cell(i, 0), img_info.get('label', '无描述'), align='CENTER')
            pic_cell = nested_table.cell(i, 1)
            pic_cell.text = ''
            p_pic = pic_cell.paragraphs[0]
            p_pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                p_pic.add_run().add_picture(img_info['path'], width=Inches(3.5))
            except Exception as e:
                p_pic.add_run(f'图片加载失败: {e}')
    else:
        set_cell_properties(photo_cell_merged, '无照片', align='CENTER')
    row_idx += 1

    # [cite_start]案例附件 [cite: 2]
    add_merged_row('案例附件', 'attachments')

    # 保存文档
    output_filename = f"report_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    output_path = os.path.join('static', 'reports', output_filename)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    document.save(output_path)

    return output_path